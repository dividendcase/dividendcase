"""Exchange rates, conversion into the home currency, and the first-run setup."""
from datetime import date

import pytest

from conftest import FAKE_RATES_BEFORE_2022, FAKE_RATES_FROM_2022, wait_until_idle
from dividendcase.services import fx


# ── parsing and converting (no database) ─────────────────────────────────────
def test_parse_ecb_history_csv():
    text = "Date,USD,JPY,CYP,GBP,\n2026-09-28,1.1378,178.5,N/A,0.85785,\n2026-09-25,1.1301,177.2,N/A,0.8561,\n"
    days = fx.parse_history_csv(text)
    assert days[date(2026, 9, 28)] == {"USD": 1.1378, "JPY": 178.5, "GBP": 0.85785}
    assert "CYP" not in days[date(2026, 9, 25)]  # N/A: no rate that day


def test_parse_ecb_xml():
    text = """<?xml version="1.0" encoding="UTF-8"?>
<gesmes:Envelope xmlns:gesmes="http://www.gesmes.org/xml/2002-08-01" xmlns="http://www.ecb.int/vocabulary/2002-08-01/eurofxref">
  <Cube>
    <Cube time='2026-09-28'><Cube currency='USD' rate='1.1378'/><Cube currency='INR' rate='109.2095'/></Cube>
    <Cube time='2026-09-25'><Cube currency='USD' rate='1.1301'/></Cube>
  </Cube>
</gesmes:Envelope>"""
    days = fx.parse_ecb_xml(text)
    assert days == {date(2026, 9, 28): {"USD": 1.1378, "INR": 109.2095}, date(2026, 9, 25): {"USD": 1.1301}}


def test_convert_uses_the_rate_on_the_day():
    table = fx.RateTable({
        date(2021, 12, 31): {"USD": 1.25, "GBP": 0.90},
        date(2022, 1, 3): {"USD": 1.10, "GBP": 0.85},  # a Monday: no rates at the weekend
    })
    assert fx.convert(110, "USD", "EUR", table) == pytest.approx(100.0)  # latest
    assert fx.convert(125, "USD", "EUR", table, date(2022, 1, 1)) == pytest.approx(100.0)  # Saturday → Friday
    assert fx.convert(100, "EUR", "GBP", table) == pytest.approx(85.0)
    assert fx.convert(1100, "USD", "GBP", table) == pytest.approx(850.0)  # cross rate through the euro
    assert fx.convert(8500, "GBp", "EUR", table) == pytest.approx(100.0)  # pence
    assert fx.convert(42, "SEK", "SEK", table) == 42  # same currency needs no rate
    assert fx.convert(10, "XYZ", "EUR", table) is None
    assert fx.convert(10, "USD", "EUR", table, date(1999, 1, 1)) == pytest.approx(8.0)  # before the first day


# ── rates stored by the app ──────────────────────────────────────────────────
@pytest.fixture
def rates(client):
    """Fetch the (fake) ECB history into the database."""
    stored = client.portal.call(fx.refresh_rates)
    assert stored > 0
    return stored


def test_rates_are_stored_once_and_served(client, rates):
    assert client.portal.call(fx.refresh_rates) == 0  # already up to date
    latest = client.get("/api/v1/fx/rates").json()
    assert latest["base"] == "EUR" and latest["rates"]["USD"] == FAKE_RATES_FROM_2022["USD"]
    assert latest["rates"]["EUR"] == 1.0
    old = client.get("/api/v1/fx/rates", params={"date": "2019-06-01"}).json()  # a Saturday
    assert old["date"] == "2019-05-31" and old["rates"]["USD"] == FAKE_RATES_BEFORE_2022["USD"]
    assert "INR" in client.get("/api/v1/fx/currencies").json()["currencies"]
    status = client.get("/api/v1/data/status").json()
    assert status["fx_latest_date"] and status["fx_days"] == rates


def test_no_rates_yet(client):
    assert client.get("/api/v1/fx/rates").json() == {
        "base": "EUR", "date": None, "rates": {}, "source": fx.SOURCE,
    }


# ── the portfolio in one currency ────────────────────────────────────────────
def test_analysis_converts_each_amount_at_its_own_date(client, market, rates):
    market.add("TESTU", amount=0.50, price=50.0, years=6)  # USD, quarterly
    market.add("TESTG", currency="GBp", exchange="LSE", amount=20.0, price=2000.0, years=6)
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 10, "purchase_price": 40.0,
                                           "purchase_date": "2021-03-01", "purchase_currency": "USD"})
    client.post("/api/v1/portfolio", json={"ticker": "TESTG", "quantity": 10, "purchase_price": 1800.0,
                                           "purchase_date": "2023-03-01", "purchase_currency": "GBp"})

    mixed = client.get("/api/v1/portfolio/analysis").json()
    assert mixed["converted"] is False
    assert mixed["total_initial_investment"] == pytest.approx(400 + 18000)  # dollars plus pence: meaningless

    eur = client.get("/api/v1/portfolio/analysis", params={"currency": "EUR"}).json()
    assert eur["converted"] is True and eur["currency"] == "EUR"
    # $400 at the 2021 rate (1.25) + 18,000p = £180 at the 2023 rate (0.85)
    assert eur["total_initial_investment"] == pytest.approx(400 / 1.25 + 180 / 0.85, rel=1e-3)
    last = eur["data_points"][-1]
    # Today's value at today's rates: 10 × $50 and 10 × 2000p
    assert last["stock_values"]["TESTU"] == pytest.approx(500 / 1.10, rel=1e-3)
    assert last["stock_values"]["TESTG"] == pytest.approx(200 / 0.85, rel=1e-3)
    assert eur["unconverted_currencies"] == []


def test_analysis_refuses_a_currency_without_rates(client, market, rates):
    market.add("TESTU")
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 1, "purchase_date": "2022-06-01"})
    r = client.get("/api/v1/portfolio/analysis", params={"currency": "XYZ"})
    assert r.status_code == 422


# ── first-run setup ──────────────────────────────────────────────────────────
def test_first_run_setup(client, rates):
    prefs = client.get("/api/v1/user/preferences").json()
    assert prefs["setup_completed_at"] is None  # the app shows the setup screen

    r = client.patch("/api/v1/user/preferences", json={
        "home_currency": "EUR", "tax_residence": "IE",
        "screener_markets": ["ISEQ20", "FTSE100"], "complete_setup": True,
    })
    assert r.status_code == 200, r.text
    prefs = r.json()
    assert prefs["home_currency"] == "EUR" and prefs["tax_residence"] == "IE"
    assert prefs["screener_markets"] == ["ISEQ20", "FTSE100"]
    assert prefs["setup_completed_at"] is not None

    # null brings back every market
    assert client.patch("/api/v1/user/preferences", json={"screener_markets": None}).json()["screener_markets"] is None


def test_setup_validates_its_answers(client, rates):
    patch = lambda body: client.patch("/api/v1/user/preferences", json=body)  # noqa: E731
    assert patch({"home_currency": "XYZ"}).status_code == 422  # no rates for it
    assert patch({"home_currency": "eur"}).status_code == 422
    assert patch({"tax_residence": "Ireland"}).status_code == 422
    assert patch({"screener_markets": ["MOON"]}).status_code == 422
    assert patch({"home_currency": "INR"}).status_code == 200


def test_markets_list(client):
    markets = client.get("/api/v1/data/markets").json()["markets"]
    assert {m["key"] for m in markets} == {"SP500", "NIFTY50", "TSX60", "ASX200", "FTSE100", "ISEQ20", "HIGH_YIELD"}
    assert all(m["stocks"] > 0 for m in markets)


def test_screener_fetches_only_the_chosen_markets(client, market):
    market.add("TESTNY")  # NYSE → S&P 500 in the fake universe
    market.add("TESTLN", currency="GBp", exchange="LSE")
    client.patch("/api/v1/user/preferences", json={"screener_markets": ["FTSE100"]})
    market.calls.clear()

    london = {t for t, (stock, _) in market.stocks.items() if stock["exchange"] == "LSE"}
    assert client.post("/api/v1/data/refresh", json={"scope": "screener"}).json()["queued"] == len(london)
    wait_until_idle(client)
    assert set(market.calls) == london and "TESTLN" in london and "TESTNY" not in market.calls


def test_cost_in_one_currency_uses_purchase_date_rates(client, market, rates):
    market.add("TESTU")
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 10, "purchase_price": 40.0,
                                           "purchase_date": "2021-03-01", "purchase_currency": "USD"})
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 5, "purchase_price": 60.0,
                                           "purchase_date": "2023-03-01", "purchase_currency": "USD"})
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 1, "purchase_date": "2024-03-01"})
    cost = client.get("/api/v1/portfolio/cost", params={"currency": "EUR"}).json()
    assert cost["total"] == pytest.approx(400 / 1.25 + 300 / 1.10, abs=0.01)
    assert cost["lots_without_price"] == 1 and cost["unconverted_currencies"] == []


def test_scheduled_screener_waits_for_setup(client, market):
    from dividendcase.services.refresh import queue_screener

    market.add("TESTNY")
    assert client.portal.call(lambda: queue_screener(scheduled=True)) == 0  # setup not finished
    client.patch("/api/v1/user/preferences", json={"complete_setup": True})
    assert client.portal.call(lambda: queue_screener(scheduled=True)) > 0
    wait_until_idle(client)


def test_excel_report_with_and_without_a_home_currency(client, market, rates):
    market.add("TESTU")
    market.add("TESTG", currency="GBp", exchange="LSE", amount=20.0, price=2000.0)
    client.post("/api/v1/portfolio", json={"ticker": "TESTU", "quantity": 10, "purchase_price": 40.0, "purchase_date": "2022-03-01"})
    client.post("/api/v1/portfolio", json={"ticker": "TESTG", "quantity": 10, "purchase_price": 1800.0,
                                           "purchase_date": "2023-03-01", "purchase_currency": "GBp"})
    plain = client.get("/api/v1/excel/report")
    assert plain.status_code == 200 and plain.content[:2] == b"PK"

    client.patch("/api/v1/user/preferences", json={"home_currency": "EUR"})
    converted = client.get("/api/v1/excel/report")
    assert converted.status_code == 200 and converted.content[:2] == b"PK"

    import io
    from openpyxl import load_workbook

    summary = load_workbook(io.BytesIO(converted.content)).worksheets[0]
    values = [cell for row in summary.iter_rows(values_only=True) for cell in row if cell is not None]
    assert "EUR" in values  # the report says which currency its totals are in
