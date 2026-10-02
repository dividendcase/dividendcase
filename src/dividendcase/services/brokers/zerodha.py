"""Zerodha's tradebook, as downloaded from Console (Reports → Tradebook), CSV or Excel.

Columns, as Zerodha publishes them: symbol, isin, trade_date, exchange, segment, series, trade_type,
auction, quantity, price, trade_id, order_id, order_execution_time. The Excel download puts a few
lines about the account above the table and title-cases the headers ("Trade Date"), so the header
row is found by its names rather than its position.

Only equity trades are read (segment EQ); derivatives, currency and commodity rows are counted and
left out. NSE symbols become Yahoo tickers ending .NS, BSE ones .BO.
"""
import csv
import io
import re
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from typing import Optional

from dividendcase.services.brokers import Trade


class TradebookError(ValueError):
    """The file isn't a Zerodha tradebook the app can read"""


REQUIRED = {"symbol", "isin", "trade_date", "trade_type", "quantity", "price"}
SUFFIX = {"NSE": ".NS", "BSE": ".BO"}
# Series codes Zerodha can append to a symbol, as in "IDEA-BE"
SERIES_SUFFIX = re.compile(r"-(EQ|BE|BZ|SM|ST|GB|GS|N[0-9A-Z]|X[0-9A-Z])$")


@dataclass
class Tradebook:
    trades: list[Trade] = field(default_factory=list)
    not_equity: int = 0  # rows for derivatives, currencies or commodities
    unreadable: list[str] = field(default_factory=list)  # one line per row that couldn't be read
    first: Optional[date] = None
    last: Optional[date] = None


def _key(header) -> str:
    return re.sub(r"[^a-z]+", "_", str(header or "").strip().lower()).strip("_")


def _rows(content: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise TradebookError("Couldn't open this Excel file.") from e
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return rows
    text = content.decode("utf-8-sig", errors="replace")
    return list(csv.reader(io.StringIO(text)))


def _date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text[:11].strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unreadable date {text!r}")


def _number(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "").strip())


def _when(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            continue
    return None


def ticker_for(symbol: str, exchange: str) -> str:
    base = SERIES_SUFFIX.sub("", symbol.strip().upper())
    return base + SUFFIX.get(exchange.strip().upper(), ".NS")


def parse(content: bytes, filename: str) -> Tradebook:
    rows = _rows(content, filename)
    header_at = next(
        (i for i, row in enumerate(rows[:40]) if REQUIRED <= {_key(c) for c in row if c is not None}),
        None,
    )
    if header_at is None:
        raise TradebookError(
            "This doesn't look like a Zerodha tradebook: it needs columns such as Symbol, ISIN, "
            "Trade Date, Trade Type, Quantity and Price."
        )
    columns = {_key(c): i for i, c in enumerate(rows[header_at]) if c is not None}

    book = Tradebook()
    for number, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        if not row or all(c in (None, "") for c in row):
            continue

        def get(name: str):
            i = columns.get(name)
            return row[i] if i is not None and i < len(row) else None

        segment = str(get("segment") or "EQ").strip().upper()
        if segment not in ("EQ", "EQUITY"):
            book.not_equity += 1
            continue
        try:
            side = str(get("trade_type")).strip().lower()
            if side not in ("buy", "sell"):
                raise ValueError(f"trade type {side!r}")
            symbol = str(get("symbol")).strip()
            isin = str(get("isin")).strip().upper()
            if not symbol or not isin:
                raise ValueError("no symbol or ISIN")
            exchange = str(get("exchange") or "NSE").strip().upper()
            trade_date = _date(get("trade_date"))
            quantity = _number(get("quantity"))
            price = _number(get("price"))
            if quantity <= 0 or price <= 0:
                raise ValueError("quantity and price must be above zero")
        except (TypeError, ValueError) as e:
            book.unreadable.append(f"Row {number}: {e}")
            continue
        trade_id = str(get("trade_id") or "").strip() or f"{isin}-{trade_date}-{side}-{quantity}-{price}-{number}"
        book.trades.append(Trade(
            isin=isin,
            symbol=symbol,
            ticker=ticker_for(symbol, exchange),
            side=side,
            trade_date=trade_date,
            quantity=quantity,
            price=price,
            currency="INR",
            trade_id=trade_id,
            executed_at=_when(get("order_execution_time")),
        ))
        book.first = min(book.first or trade_date, trade_date)
        book.last = max(book.last or trade_date, trade_date)

    # A stock bought on both exchanges is one holding: use its NSE ticker wherever it has one
    nse = {t.isin: t.ticker for t in book.trades if t.ticker.endswith(".NS")}
    book.trades = [replace(t, ticker=nse[t.isin]) if t.isin in nse else t for t in book.trades]
    return book
