"""Reading brokers' files: spreadsheet rows, header names, dates and numbers as brokers write them."""
import csv
import io
import re
from datetime import date, datetime
from typing import Optional


class BrokerFileError(ValueError):
    """The file isn't one the app can read; the message says why, for the user"""


def key(header) -> str:
    """A header as a plain name: "Trade Date" → "trade_date", "Avg Trading\\nPrice" → "avg_trading_price" """
    return re.sub(r"[^a-z]+", "_", str(header or "").strip().lower()).strip("_")


def rows(content: bytes, filename: str) -> list[list]:
    """Every row of a .csv file, or of an .xlsx file's first sheet"""
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise BrokerFileError("Couldn't open this Excel file.") from e
        found = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return found
    text = content.decode("utf-8-sig", errors="replace")
    return list(csv.reader(io.StringIO(text)))


DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y")


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text[:11].strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unreadable date {text!r}")


def parse_number(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "").strip())


def parse_datetime(value) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            continue
    return None
